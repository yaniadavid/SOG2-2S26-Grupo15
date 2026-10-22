"""Reabre Manual_1.docx y comprueba su estructura y la ausencia de IP, contrasenas y emojis.

Uso: python verificar_docx.py [ruta_docx]
Termina con codigo 1 si encuentra IP, patrones de contrasena o emojis.
"""
import re
import sys
from pathlib import Path

from docx import Document

CARPETA = Path(__file__).resolve().parent

# Se excluye 0.12.6.1, que es la version de wkhtmltopdf y no una direccion.
PATRON_IP = re.compile(r"\b(?!0\.12\.6\.1\b)(?:\d{1,3}\.){3}\d{1,3}\b")
PATRON_CLAVE = re.compile(
    r"(?i)(passw(?:or)?d|passwd|contrase[nñ]a|clave|secret|token)\s*[:=]\s*[^\s<\[(]{2,}"
)
PATRON_URL_CREDENCIALES = re.compile(r"://[^/\s:@]+:[^/\s@]+@")
PATRON_EMOJI = re.compile(
    "[\U0001F000-\U0001FAFF\u2600-\u27bf\u2b00-\u2bff\u2300-\u23ff\ufe0f\u200d\u2190-\u21ff]"
)
PATRON_MARCADOR = re.compile(r"MARCADOR DE IMAGEN - Captura (\d+) pendiente")


def todo_el_texto(documento) -> list:
    textos = [p.text for p in documento.paragraphs]
    for tabla in documento.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                textos.extend(p.text for p in celda.paragraphs)
    return textos


def main(argv: list) -> int:
    ruta = Path(argv[0]) if argv else CARPETA / "Manual_1.docx"
    documento = Document(str(ruta))
    titulos = {1: 0, 2: 0, 3: 0}
    for parrafo in documento.paragraphs:
        estilo = parrafo.style.name
        if estilo.startswith("Heading "):
            nivel = int(estilo.split()[-1])
            if nivel in titulos:
                titulos[nivel] += 1
    textos = todo_el_texto(documento)
    marcadores = sum(1 for t in textos if PATRON_MARCADOR.search(t))
    imagenes = len(documento.inline_shapes)
    campo_toc = "TOC" in documento.element.xml and "fldChar" in documento.element.xml

    print(f"Archivo: {ruta.name} ({ruta.stat().st_size} bytes)")
    print(f"Titulos nivel 1/2/3: {titulos[1]}/{titulos[2]}/{titulos[3]}")
    print(f"Tablas: {len(documento.tables)}; imagenes: {imagenes}; marcadores de captura: {marcadores}")
    print(f"Campo TOC presente: {campo_toc}")

    hallazgos = []
    for texto in textos:
        for nombre, patron in (("IP", PATRON_IP), ("clave", PATRON_CLAVE),
                               ("url con credenciales", PATRON_URL_CREDENCIALES),
                               ("emoji", PATRON_EMOJI)):
            for coincidencia in patron.finditer(texto):
                hallazgos.append((nombre, coincidencia.group(0)))
    # Contenido fuente
    fuente = (CARPETA / "contenido.json").read_text(encoding="utf-8")
    for nombre, patron in (("IP", PATRON_IP), ("clave", PATRON_CLAVE), ("emoji", PATRON_EMOJI)):
        for coincidencia in patron.finditer(fuente):
            hallazgos.append((nombre + " (contenido.json)", coincidencia.group(0)))
    print(f"Hallazgos de IP, clave o emoji: {len(hallazgos)}")
    for nombre, valor in hallazgos:
        print(f"  {nombre}: {valor}")
    return 1 if hallazgos else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
