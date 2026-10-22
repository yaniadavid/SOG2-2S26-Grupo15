"""Genera Manual_1.docx a partir de contenido.json.

La seccion 2 se lee en tiempo de ejecucion de docs/manual1/funcionamiento_modulos.txt (o de la ruta
indicada con --funcionamiento); si no existe, se usa un texto provisional marcado.
Las capturas se buscan en docs/manual1/capturas/ como cap_NN.png (o cap_NN_tema.png); si no existen se
deja un recuadro de marcador visible.

Uso: python generar_docx.py [--funcionamiento RUTA] [--capturas CARPETA] [--salida ARCHIVO]
"""
import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CARPETA = Path(__file__).resolve().parent
CARPETA_DOC = CARPETA.parent
CONTENIDO = CARPETA / "contenido.json"
FUNCIONAMIENTO = CARPETA_DOC / "funcionamiento_modulos.txt"
CAPTURAS = CARPETA_DOC / "capturas"
SALIDA = CARPETA / "Manual_1.docx"

COLOR_ENCABEZADO = "1F3864"
COLOR_FILA_ALTERNA = "EEF2F8"
COLOR_CODIGO = "F2F2F2"
COLOR_MARCADOR = "FFF2CC"
ANCHO_IMAGEN_CM = 15.0
FUENTE = "Calibri"
FUENTE_CODIGO = "Consolas"

RE_CAPTURA_SOLA = re.compile(r"^\[?\s*Captura\s+(\d+)\s*\]?\s*[:.]?\s*(.*)$", re.IGNORECASE)
RE_ENCABEZADO = re.compile(r"^(\d+(?:\.\d+)+)\.?\s+(.+)$")
RE_SEPARADOR = re.compile(r"^[=\-_*#~\s]{3,}$")

TEXTO_PROVISIONAL = [
    "PROVISIONAL: el archivo docs/manual1/funcionamiento_modulos.txt todavía no existe. Este texto se "
    "reemplaza automáticamente al volver a generar el manual cuando el archivo esté disponible.",
]
SUBTITULOS_PROVISIONALES = [
    "2.1 Ventas", "2.2 Compras", "2.3 Facturación", "2.4 Empleados",
]

# ----------------------------------------------------------------------- utilidades XML


def sombrear(celda, color: str) -> None:
    propiedades = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    propiedades.append(sombra)


def agregar_campo(parrafo, instruccion: str, texto_visible: str = "") -> None:
    """Inserta un campo de Word (complejo) en un parrafo."""
    def elemento_run():
        corrida = parrafo.add_run()
        return corrida, corrida._r

    corrida, r = elemento_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    r.append(inicio)
    corrida, r = elemento_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruccion} "
    r.append(instr)
    corrida, r = elemento_run()
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    r.append(separador)
    parrafo.add_run(texto_visible)
    corrida, r = elemento_run()
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    r.append(fin)


def pedir_actualizar_campos(documento) -> None:
    ajustes = documento.settings.element
    actualizar = OxmlElement("w:updateFields")
    actualizar.set(qn("w:val"), "true")
    ajustes.append(actualizar)


def repetir_encabezado(fila) -> None:
    propiedades = fila._tr.get_or_add_trPr()
    marca = OxmlElement("w:tblHeader")
    marca.set(qn("w:val"), "true")
    propiedades.append(marca)


def evitar_corte_fila(fila) -> None:
    propiedades = fila._tr.get_or_add_trPr()
    marca = OxmlElement("w:cantSplit")
    marca.set(qn("w:val"), "true")
    propiedades.append(marca)


# ------------------------------------------------------------------------- estilos


def configurar_estilos(documento) -> None:
    normal = documento.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    lenguaje = OxmlElement("w:lang")
    lenguaje.set(qn("w:val"), "es-GT")
    normal.element.rPr.append(lenguaje)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for nombre, tam in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        estilo = documento.styles[nombre]
        estilo.font.name = FUENTE
        estilo.font.size = Pt(tam)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor.from_string(COLOR_ENCABEZADO)
        estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        estilo.paragraph_format.space_before = Pt(14 if nombre != "Heading 3" else 10)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.keep_with_next = True


def configurar_pagina(documento) -> None:
    seccion = documento.sections[0]
    seccion.page_width, seccion.page_height = Cm(21.59), Cm(27.94)
    seccion.left_margin = seccion.right_margin = Cm(2.5)
    seccion.top_margin = seccion.bottom_margin = Cm(2.3)
    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.add_run("Manual 1 - Grupo 15 - Página ")
    agregar_campo(pie, "PAGE", "1")
    seccion.different_first_page_header_footer = True


# ----------------------------------------------------------------------- elementos


def parrafo_texto(documento, texto: str, estilo: str = None, negrita: bool = False,
                  cursiva: bool = False, alineacion=None):
    parrafo = documento.add_paragraph(style=estilo)
    corrida = parrafo.add_run(texto)
    corrida.bold = negrita
    corrida.italic = cursiva
    if alineacion is not None:
        parrafo.alignment = alineacion
    return parrafo


def agregar_lista(documento, items: list, numerada: bool) -> None:
    for numero, item in enumerate(items, start=1):
        if numerada:
            # Numeracion manual: los estilos de lista de Word encadenan la cuenta entre listas.
            parrafo = documento.add_paragraph(f"{numero}.\t{item}")
            parrafo.paragraph_format.left_indent = Cm(1.0)
            parrafo.paragraph_format.first_line_indent = Cm(-0.7)
            parrafo.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
        else:
            documento.add_paragraph(item, style="List Bullet")


def agregar_tabla(documento, columnas: list, filas: list) -> None:
    tabla = documento.add_table(rows=1, cols=len(columnas))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    encabezado = tabla.rows[0]
    repetir_encabezado(encabezado)
    for celda, texto in zip(encabezado.cells, columnas):
        celda.text = ""
        corrida = celda.paragraphs[0].add_run(texto)
        corrida.bold = True
        corrida.font.size = Pt(10)
        corrida.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(celda, COLOR_ENCABEZADO)
    for indice, datos in enumerate(filas):
        fila = tabla.add_row()
        evitar_corte_fila(fila)
        for celda, texto in zip(fila.cells, datos):
            celda.text = ""
            corrida = celda.paragraphs[0].add_run(str(texto))
            corrida.font.size = Pt(10)
            if indice % 2 == 1:
                sombrear(celda, COLOR_FILA_ALTERNA)
    documento.add_paragraph()


def agregar_codigo(documento, texto: str) -> None:
    tabla = documento.add_table(rows=1, cols=1)
    tabla.style = "Table Grid"
    celda = tabla.rows[0].cells[0]
    sombrear(celda, COLOR_CODIGO)
    celda.text = ""
    lineas = texto.split("\n")
    for numero, linea in enumerate(lineas):
        parrafo = celda.paragraphs[0] if numero == 0 else celda.add_paragraph()
        parrafo.paragraph_format.space_after = Pt(0)
        parrafo.paragraph_format.line_spacing = 1.0
        corrida = parrafo.add_run(linea if linea else " ")
        corrida.font.name = FUENTE_CODIGO
        corrida.font.size = Pt(9)
        corrida._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FUENTE_CODIGO)
    documento.add_paragraph()


def agregar_nota(documento, texto: str, etiqueta: str = "Nota") -> None:
    tabla = documento.add_table(rows=1, cols=1)
    tabla.style = "Table Grid"
    celda = tabla.rows[0].cells[0]
    sombrear(celda, COLOR_MARCADOR)
    celda.text = ""
    parrafo = celda.paragraphs[0]
    corrida = parrafo.add_run(f"{etiqueta}: ")
    corrida.bold = True
    parrafo.add_run(texto)
    documento.add_paragraph()


def buscar_captura(carpeta: Path, numero: int):
    exacta = carpeta / f"cap_{numero:02d}.png"
    if exacta.is_file():
        return exacta
    for candidata in sorted(carpeta.glob(f"cap_{numero:02d}_*.png")):
        return candidata
    return None


def agregar_captura(documento, numero: int, leyenda: str, carpeta: Path, estadisticas: dict) -> None:
    ruta = buscar_captura(carpeta, numero)
    if ruta is not None:
        parrafo = documento.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parrafo.paragraph_format.keep_with_next = True
        parrafo.add_run().add_picture(str(ruta), width=Cm(ANCHO_IMAGEN_CM))
        estadisticas["imagenes"] += 1
        pie = documento.add_paragraph()
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        corrida = pie.add_run(f"Captura {numero}. {leyenda}")
        corrida.italic = True
        corrida.font.size = Pt(9)
        return
    tabla = documento.add_table(rows=1, cols=1)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = tabla.rows[0].cells[0]
    sombrear(celda, COLOR_MARCADOR)
    celda.text = ""
    parrafo = celda.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corrida = parrafo.add_run(f"MARCADOR DE IMAGEN - Captura {numero} pendiente")
    corrida.bold = True
    segundo = celda.add_paragraph()
    segundo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    segundo.add_run(leyenda)
    tercero = celda.add_paragraph()
    tercero.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tercero.add_run(f"Archivo esperado: docs/manual1/capturas/cap_{numero:02d}.png")
    r.italic = True
    r.font.size = Pt(9)
    estadisticas["marcadores"] += 1
    documento.add_paragraph()


# ------------------------------------------------------------------------ portada


def agregar_portada(documento, datos: dict) -> None:
    for _ in range(3):
        documento.add_paragraph()
    parrafo_texto(documento, datos["universidad"], negrita=True,
                  alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3):
        documento.add_paragraph()
    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corrida = titulo.add_run(datos["titulo"])
    corrida.bold = True
    corrida.font.size = Pt(26)
    corrida.font.color.rgb = RGBColor.from_string(COLOR_ENCABEZADO)
    sub = documento.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corrida = sub.add_run(datos["subtitulo"])
    corrida.font.size = Pt(16)
    for _ in range(4):
        documento.add_paragraph()
    for etiqueta, clave in (("Curso", "curso"), ("Grupo", "grupo"),
                            ("Integrante", "integrante"), ("Fecha", "fecha")):
        linea = documento.add_paragraph()
        linea.alignment = WD_ALIGN_PARAGRAPH.CENTER
        e = linea.add_run(f"{etiqueta}: ")
        e.bold = True
        linea.add_run(datos[clave])
    documento.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def agregar_tabla_contenido(documento) -> None:
    titulo = documento.add_paragraph()
    corrida = titulo.add_run("Tabla de contenido")
    corrida.bold = True
    corrida.font.size = Pt(18)
    corrida.font.color.rgb = RGBColor.from_string(COLOR_ENCABEZADO)
    parrafo = documento.add_paragraph()
    agregar_campo(parrafo, 'TOC \\o "1-3" \\h \\z \\u',
                  "Haga clic derecho y elija Actualizar campo para generar la tabla de contenido.")
    documento.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------- seccion 2 (externa)


RE_CLAVE = re.compile(r"^[\w.]{2,30}:\s")
RE_MARCADOR = re.compile(r"^(?:[-*•]\s+|\d+\.\s+)")


def sangria(linea: str) -> int:
    return len(linea) - len(linea.lstrip(" "))


def limpiar(texto: str) -> str:
    return " ".join(texto.split())


def procesar_fragmento(fragmento: list) -> list:
    """Convierte lineas consecutivas (sangria, texto) en parrafos, etiquetas y listas."""
    salida, i = [], 0
    while i < len(fragmento):
        nivel, texto = fragmento[i]
        siguiente = fragmento[i + 1][0] if i + 1 < len(fragmento) else None
        if nivel == 0 and not RE_MARCADOR.match(texto) and siguiente and siguiente > 0 and len(texto) <= 80:
            salida.append({"t": "label", "texto": limpiar(texto)})
            i += 1
            continue
        if nivel == 0 and not RE_MARCADOR.match(texto):
            partes = []
            while i < len(fragmento) and fragmento[i][0] == 0 and not RE_MARCADOR.match(fragmento[i][1]):
                partes.append(limpiar(fragmento[i][1]))
                i += 1
            salida.append({"t": "p", "texto": " ".join(partes)})
            continue
        base, items = nivel, []
        hay_mas_sangria = any(n > base for n, _ in fragmento[i:])
        while i < len(fragmento) and fragmento[i][0] >= base:
            nivel_i, texto_i = fragmento[i]
            inicia_item = (not items or RE_MARCADOR.match(texto_i) or " -> " in texto_i
                           or RE_CLAVE.match(texto_i) or (base > 0 and hay_mas_sangria))
            if nivel_i == base and inicia_item:
                items.append(limpiar(texto_i))
            else:
                items[-1] += " " + limpiar(texto_i)
            i += 1
        numerados = all(re.match(r"^\d+\.\s", it) for it in items)
        limpios = [RE_MARCADOR.sub("", it, count=1) for it in items]
        salida.append({"t": "ol" if numerados else "ul", "items": limpios})
    return salida


def analizar_funcionamiento(texto: str) -> list:
    """Convierte el texto plano en bloques (h2, h3, label, p, ul, ol, table, code, cap)."""
    bloques, fragmento, tabla = [], [], []
    primera_vista = False  # la primera linea en mayusculas es el titulo del archivo y se omite

    def cerrar_fragmento():
        if fragmento:
            bloques.extend(procesar_fragmento(fragmento))
            fragmento.clear()

    def cerrar_tabla():
        if len(tabla) >= 2:
            columnas = tabla[0]
            filas = [f + [""] * (len(columnas) - len(f)) for f in tabla[1:]]
            bloques.append({"t": "table", "cols": columnas, "filas": filas})
        elif tabla:
            bloques.append({"t": "p", "texto": " | ".join(tabla[0])})
        tabla.clear()

    for linea_cruda in texto.splitlines():
        linea = linea_cruda.rstrip().expandtabs(4)
        limpia = linea.strip()
        if not limpia or RE_SEPARADOR.match(limpia):
            cerrar_fragmento()
            cerrar_tabla()
            continue
        if limpia.upper().startswith(("SECCION", "SECCIÓN")):
            cerrar_fragmento()
            cerrar_tabla()
            continue
        if "|" in limpia and sangria(linea) == 0:
            cerrar_fragmento()
            tabla.append([c.strip() for c in limpia.strip("|").split("|")])
            continue
        cerrar_tabla()
        captura = RE_CAPTURA_SOLA.match(limpia)
        if captura and not captura.group(2):
            cerrar_fragmento()
            bloques.append({"t": "cap", "n": int(captura.group(1))})
            continue
        encabezado = RE_ENCABEZADO.match(limpia)
        if encabezado and len(limpia) <= 110 and sangria(linea) == 0:
            cerrar_fragmento()
            nivel = 2 if encabezado.group(1).count(".") == 1 else 3
            bloques.append({"t": "h2" if nivel == 2 else "h3", "texto": limpia})
            continue
        if limpia.startswith("$ "):
            cerrar_fragmento()
            bloques.append({"t": "code", "texto": limpia[2:]})
            continue
        if sangria(linea) == 0 and limpia.isupper():
            cerrar_fragmento()
            if bloques or primera_vista:
                bloques.append({"t": "label", "texto": limpia})
            primera_vista = True
            continue
        primera_vista = True
        fragmento.append((sangria(linea), limpia))
    cerrar_fragmento()
    cerrar_tabla()
    return bloques


def bloques_seccion2(ruta: Path):
    if ruta.is_file():
        texto = ruta.read_text(encoding="utf-8")
        return analizar_funcionamiento(texto), True
    bloques = [{"t": "note", "texto": TEXTO_PROVISIONAL[0], "etiqueta": "PROVISIONAL"}]
    for subtitulo in SUBTITULOS_PROVISIONALES:
        bloques.append({"t": "h2", "texto": subtitulo})
        bloques.append({"t": "p", "texto": "Contenido pendiente de recibir."})
    return bloques, False


# --------------------------------------------------------------------- render


def renderizar(documento, bloques: list, capturas: dict, carpeta_capturas: Path,
               estadisticas: dict) -> None:
    for bloque in bloques:
        tipo = bloque["t"]
        if tipo == "p":
            parrafo_texto(documento, bloque["texto"])
        elif tipo == "label":
            parrafo_texto(documento, bloque["texto"], negrita=True).paragraph_format.keep_with_next = True
        elif tipo == "h2":
            documento.add_heading(bloque["texto"], level=2)
        elif tipo == "h3":
            documento.add_heading(bloque["texto"], level=3)
        elif tipo == "ul":
            agregar_lista(documento, bloque["items"], numerada=False)
        elif tipo == "ol":
            agregar_lista(documento, bloque["items"], numerada=True)
        elif tipo == "table":
            agregar_tabla(documento, bloque["cols"], bloque["filas"])
            estadisticas["tablas_contenido"] += 1
        elif tipo == "code":
            agregar_codigo(documento, bloque["texto"])
        elif tipo == "note":
            agregar_nota(documento, bloque["texto"], bloque.get("etiqueta", "Nota"))
        elif tipo == "pend":
            agregar_nota(documento, bloque["texto"], "Pendiente")
        elif tipo == "cap":
            numero = bloque["n"]
            leyenda = bloque.get("pie") or capturas.get(str(numero), "Evidencia del sistema")
            agregar_captura(documento, numero, leyenda, carpeta_capturas, estadisticas)
        else:
            raise ValueError(f"Tipo de bloque desconocido: {tipo}")


def renderizar_subs(documento, seccion: dict, nivel: int, contexto: dict) -> None:
    renderizar(documento, seccion.get("bloques", []), *contexto["args"])
    for sub in seccion.get("subs", []):
        documento.add_heading(sub["titulo"], level=min(nivel + 1, 3))
        renderizar_subs(documento, sub, nivel + 1, contexto)


def construir(contenido: dict, ruta_funcionamiento: Path, carpeta_capturas: Path,
              ruta_salida: Path) -> dict:
    documento = Document()
    configurar_estilos(documento)
    configurar_pagina(documento)
    agregar_portada(documento, contenido["portada"])
    agregar_tabla_contenido(documento)
    estadisticas = {"imagenes": 0, "marcadores": 0, "tablas_contenido": 0, "seccion2_externa": False}
    contexto = {"args": (contenido["capturas"], carpeta_capturas, estadisticas)}

    for indice, seccion in enumerate(contenido["secciones"]):
        titulo = documento.add_heading(seccion["titulo"], level=1)
        if indice > 0:
            titulo.paragraph_format.page_break_before = True
        if seccion.get("fuente_externa"):
            bloques, existe = bloques_seccion2(ruta_funcionamiento)
            estadisticas["seccion2_externa"] = existe
            renderizar(documento, bloques, *contexto["args"])
        else:
            renderizar_subs(documento, seccion, 1, contexto)
    pedir_actualizar_campos(documento)
    documento.save(str(ruta_salida))
    return estadisticas


def main(argv: list) -> int:
    analizador = argparse.ArgumentParser(description="Genera Manual_1.docx")
    analizador.add_argument("--contenido", type=Path, default=CONTENIDO)
    analizador.add_argument("--funcionamiento", type=Path, default=FUNCIONAMIENTO)
    analizador.add_argument("--capturas", type=Path, default=CAPTURAS)
    analizador.add_argument("--salida", type=Path, default=SALIDA)
    args = analizador.parse_args(argv)
    contenido = json.loads(args.contenido.read_text(encoding="utf-8"))
    estadisticas = construir(contenido, args.funcionamiento, args.capturas, args.salida)
    print(f"Generado: {args.salida}")
    print(f"Imagenes insertadas: {estadisticas['imagenes']}; marcadores: {estadisticas['marcadores']}; "
          f"tablas de contenido: {estadisticas['tablas_contenido']}; "
          f"seccion 2 desde archivo: {estadisticas['seccion2_externa']}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
